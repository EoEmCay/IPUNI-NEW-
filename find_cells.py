import openpyxl

wb = openpyxl.load_workbook(r'd:\DIA+\DIA+_Mô hình tài chính - đã hiệu chỉnh.xlsx', data_only=False)

# Inspect DIA_Chi Phí
ws_cp = wb['DIA_Chi Phí']
print("--- DIA_Chi Phí ---")
for row in ws_cp.iter_rows(min_row=110, max_row=130):
    for cell in row:
        if cell.value and isinstance(cell.value, str) and 'Marketing' in cell.value:
            print(f"Found '{cell.value}' at {cell.coordinate}")
            # print adjacent cells
            print(f"Row {cell.row} values: {[ws_cp.cell(row=cell.row, column=c).value for c in range(1, 10)]}")

# Inspect DIA_Tóm Tắt
ws_tt = wb['DIA_Tóm Tắt']
print("--- DIA_Tóm Tắt ---")
for row in ws_tt.iter_rows(min_row=1, max_row=20):
    for cell in row:
        if cell.value == 'IRR':
            print(f"Found 'IRR' at {cell.coordinate}")
            print(f"Adjacent cell value: {ws_tt.cell(row=cell.row, column=cell.column+1).value}")

# Inspect DIA_Hiệu Quả TC
ws_hq = wb['DIA_Hiệu Quả TC']
print("--- DIA_Hiệu Quả TC ---")
for row in ws_hq.iter_rows(min_row=45, max_row=60):
    for cell in row:
        if cell.value == 'Dòng tiền đầu tư ban đầu':
            print(f"Found 'Dòng tiền đầu tư ban đầu' at {cell.coordinate}")
            print(f"Row {cell.row} values: {[ws_hq.cell(row=cell.row, column=c).value for c in range(1, 8)]}")
        if cell.value == 'Dòng tiền từ hoạt động kinh doanh':
            print(f"Found 'Dòng tiền từ hoạt động kinh doanh' at {cell.coordinate}")
            print(f"Row {cell.row} values: {[ws_hq.cell(row=cell.row, column=c).value for c in range(1, 8)]}")

